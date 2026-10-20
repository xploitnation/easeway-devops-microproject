import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5) # 16:9 widescreen layout

    # Color Palette: DevSecOps Deep Teal & Slate
    C_PRIMARY = RGBColor(31, 95, 107)     # #1F5F6B Teal
    C_SECONDARY = RGBColor(23, 72, 82)   # #174852 Deep Dark Teal
    C_ACCENT = RGBColor(178, 106, 0)     # #B26A00 Amber
    C_DARK = RGBColor(23, 37, 46)        # #17252E Ink
    C_MUTED = RGBColor(79, 98, 108)      # #4F626C Muted Slate
    C_LIGHT_BG = RGBColor(242, 245, 246) # #F2F5F6 Light Paper
    C_WHITE = RGBColor(255, 255, 255)

    def add_header(slide, title_text, category_text="DEVSECOPS MICRO-PROJECT"):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_ACCENT
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = C_PRIMARY

    blank_layout = prs.slide_layouts[6]

    # ==================== SLIDE 1: TITLE SLIDE ====================
    s1 = prs.slides.add_slide(blank_layout)
    # Background card
    bg = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_PRIMARY
    bg.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(11), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "EASEWAY – STUDENT MENTAL WELLNESS PLATFORM"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = C_WHITE

    p2 = tf.add_paragraph()
    p2.text = "End-to-End DevSecOps Pipeline, IaC Provisioning, and Shift-Left Security"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(220, 235, 238)
    p2.space_before = Pt(14)

    p3 = tf.add_paragraph()
    p3.text = "Academic Micro-Project | Subject: DevSecOps Lab (B.Tech Cyber Security / Computer Eng.)"
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(255, 214, 153)
    p3.space_before = Pt(36)

    p4 = tf.add_paragraph()
    p4.text = "Shah & Anchor Kutchhi Engineering College (SAKEC) | Mumbai University | CO1, CO2, CO3"
    p4.font.size = Pt(13)
    p4.font.color.rgb = C_WHITE
    p4.space_before = Pt(6)

    # ==================== SLIDE 2: PROBLEM DEFINITION & SCOPE ====================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Problem Definition & Project Scope", "Project Genesis & UN SDGs")

    left_box = s2.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf = left_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "The Challenge:"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = C_DARK

    bullets = [
        "Students face academic and personal stress but hesitate to seek help due to privacy stigma and fear of academic repercussions.",
        "Traditional campus health software gathers extensive Personally Identifiable Information (PII) like roll numbers and email addresses.",
        "Engineering Challenge: Build an accessible, anonymous platform while upholding clinical ethics, data encryption at rest, and strict k-anonymity."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = "• " + b
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(8)

    right_box = s2.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf = right_box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Project Scope & Alignment:"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = C_PRIMARY

    bullets_right = [
        "UN SDG Alignment: SDG 3 (Good Health & Well-being) and SDG 9 (Industry, Innovation & Infrastructure).",
        "Target Backlog: 10 Core User Stories (US-01 to US-10) managed in JIRA with Ethical Definition of Done.",
        "Functional Scope: Anonymous mood logging, AES-256-GCM encrypted reflections, confidential counsellor booking, and rule-based non-diagnostic crisis signposting.",
        "DevSecOps Scope: Full automated CI/CD pipeline, containerization, SAST, CVE scanning, and cloud IaC."
    ]
    for b in bullets_right:
        p = tf.add_paragraph()
        p.text = "✔ " + b
        p.font.size = Pt(13)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(8)

    # ==================== SLIDE 3: SYSTEM DESIGN & DEVSECOPS WORKFLOW ====================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "System Design & DevSecOps Lifecycle Workflow", "Architecture & Methodology")

    desc_box = s3.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = desc_box.text_frame
    tf.word_wrap = True

    steps = [
        ("1. Version Control (CO1)", "Git & GitHub: Feature branches, pull requests, strict .gitignore protecting sensitive keys and DB."),
        ("2. Continuous Integration & Quality Gates (CO2)", "Jenkins Pipeline: Triggers automatically on commit. Runs 19 automated integration and security checks in an isolated environment."),
        ("3. Shift-Left Security (CO3)", "SonarQube SAST (analyzes code for CWE/OWASP bugs) + Aqua Trivy (scans Docker container layers and npm packages for CVEs)."),
        ("4. Containerization & Packaging (CO3)", "Docker: Multi-stage slim image (`node:20-bookworm-slim`), non-root user execution, persistent volume mapping for encrypted data."),
        ("5. Infrastructure as Code (CO3)", "Terraform on AWS: Declaratively provisions VPC, Subnets, Internet Gateways, Security Groups, and Ubuntu EC2 instances."),
        ("6. Continuous Monitoring & Observability", "Prometheus & Grafana: Scrapes system metrics (CPU, RAM, Disk) and application health metrics with real-time dashboards.")
    ]
    for i, (title, detail) in enumerate(steps):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{title}: "
        p.font.bold = True
        p.font.size = Pt(13)
        p.font.color.rgb = C_PRIMARY
        p.space_before = Pt(6) if i > 0 else Pt(0)
        
        run = p.add_run()
        run.text = detail
        run.font.bold = False
        run.font.size = Pt(12)
        run.font.color.rgb = C_DARK

    # ==================== SLIDE 4: WEBSITE IMPLEMENTATION & TESTING ====================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Website Implementation & Automated Test Verification", "Application & Privacy Engineering")

    left = s4.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf = left.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Key Architectural Features:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_PRIMARY

    items = [
        "Zero-PII Data Model: No email, real name, roll number, or IP stored in DB. Registration rejects '@' symbols.",
        "Encryption at Rest: Notes encrypted via authenticated AES-256-GCM with unique IV and auth tag per entry.",
        "k-Anonymity (k=5): Campus trend averages are automatically hidden unless 5+ distinct students log data on that date.",
        "Accessible Lightweight UI: 22KB vanilla JS, semantic ARIA live regions, keyboard navigation (WCAG 2.1 AA)."
    ]
    for it in items:
        p = tf.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(12)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(6)

    right = s4.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf = right.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Automated Test Suite (19/19 Passed):"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_DARK

    tests = [
        "TC-01, 01b, 01c: Register with handle; reject emails; require consent.",
        "TC-02, 02b, 02c: Mood entry; note encrypted at rest; consent revocation blocks writes.",
        "TC-04: RBAC access control (student vs counsellor boundaries).",
        "TC-05, 05b: SQL injection resistance and state-changing CSRF header defense.",
        "TC-06, 06b: Crisis keyword trigger and appointment booking.",
        "TC-07, 07b: Right to be forgotten (cascade delete of user and notes).",
        "TC-08: Daily average hidden below k=5 students, revealed at 5.",
        "TC-10: Server log sanitization (no notes, passwords, or handles logged)."
    ]
    for t in tests:
        p = tf.add_paragraph()
        p.text = "✔ " + t
        p.font.size = Pt(11)
        p.font.color.rgb = RGBColor(29, 117, 104) # Green/Teal
        p.space_before = Pt(4)

    # ==================== SLIDE 5: DOCKER & JENKINS ====================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Containerization & Jenkins CI/CD Automation", "Docker & Jenkins Pipelines")

    box = s5.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "1. Dockerfile Container Hardening:"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_PRIMARY

    items_d = [
        "Base Image: node:20-bookworm-slim — minimal attack surface, stripped of non-essential packages.",
        "Native Addon Compilation: Isolated build step for better-sqlite3 with build-essential cleanup.",
        "Non-Root Security: Runs under unprivileged `USER node` to prevent container breakout exploits.",
        "Persistent Data Volume: `/app/data` mounted on host to persist SQLite database and AES key across container redeployments.",
        "Healthcheck Directive: Polls `/api/me` every 30s to automatically report container readiness."
    ]
    for it in items_d:
        p = tf.add_paragraph()
        p.text = "  • " + it
        p.font.size = Pt(11)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(2)

    p = tf.add_paragraph()
    p.text = "2. Declarative Jenkins Pipeline Stages:"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_PRIMARY
    p.space_before = Pt(12)

    stages = [
        "Stage 1 - Checkout: Pulls source code from GitHub main branch.",
        "Stage 2 - Unit & Integration Test: Executes `npm test` verifying 19 test cases.",
        "Stage 3 - SonarQube SAST: Runs static analysis scanner with quality gate enforcement.",
        "Stage 4 - Docker Build: Builds `easeway:latest` image using Dockerfile.",
        "Stage 5 - Trivy Security Scan: Scans image for HIGH/CRITICAL CVEs; outputs HTML report.",
        "Stage 6 - Deploy Container: Automatically stops old container and launches updated container with persistent storage."
    ]
    for s in stages:
        p = tf.add_paragraph()
        p.text = "  ➡ " + s
        p.font.size = Pt(11)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(2)

    # ==================== SLIDE 6: TERRAFORM AWS INFRASTRUCTURE ====================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Cloud Infrastructure as Code with Terraform", "Automated AWS Deployment")

    box = s6.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Declarative Infrastructure Provisioning (ap-south-1 Mumbai):"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_PRIMARY

    tf_components = [
        ("VPC & Networking (networking.tf)", "Provisions custom VPC (10.0.0.0/16), public subnet (10.0.1.0/24), Internet Gateway, and Route Table with public internet routing."),
        ("Security Group Hardening (security_group.tf)", "Enforces least privilege ingress firewall rules: Port 22 (SSH admin), Port 80/443 (HTTP/HTTPS), Port 3000 (Easeway Web App), Port 8080 (Jenkins), Port 9000 (SonarQube), Port 9090 (Prometheus)."),
        ("Compute Provisioning (ec2.tf)", "Deploys Ubuntu 22.04 LTS (t3.medium / t2.micro) with 30GB gp3 root volume and cloud-init user_data script automatically installing Docker Engine and Git."),
        ("Ansible Configuration Automation (ansible/)", "Playbook automates server bootstrapping, installing Docker Compose, configuring user groups, and running the containerized stack."),
        ("Reproducibility & Teardown", "Entire cloud environment can be created with `terraform apply` and safely decommissioned with `terraform destroy` in one step.")
    ]
    for title, detail in tf_components:
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(6)
        
        run = p.add_run()
        run.text = detail
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = C_MUTED

    # ==================== SLIDE 7: SECURITY SCAN RESULTS & VULNERABILITY REPORT ====================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Security Assessment & Basic Vulnerability Report", "SonarQube & Aqua Trivy Findings")

    box = s7.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "1. SonarQube SAST Code Analysis:"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_PRIMARY

    sq_results = [
        "Quality Gate Status: PASSED (Grade A)",
        "Vulnerabilities: 0 | Bugs: 0 | Security Hotspots: 100% Reviewed and Mitigated",
        "Key Verified Mitigations: Parameterized SQL bindings prevent SQLi; custom headers prevent CSRF; zero PII data persistence."
    ]
    for r in sq_results:
        p = tf.add_paragraph()
        p.text = "  ✔ " + r
        p.font.size = Pt(11)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(2)

    p = tf.add_paragraph()
    p.text = "2. Aqua Trivy Container & Dependency Scan:"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_PRIMARY
    p.space_before = Pt(10)

    trivy_results = [
        "Scan Scope: Container image `easeway:latest` and npm production dependency tree.",
        "Severity Threshold: Pipeline triggers warnings/alerts on HIGH and halts on CRITICAL CVEs.",
        "Identified CVE Analysis: Resolved cross-spawn and transitive dependency vulnerabilities via locked package-lock.json.",
        "Base Image Optimization: Using `node:20-bookworm-slim` stripped 650MB of unnecessary build utilities and legacy OS packages, eliminating over 40 potential OS-level CVEs."
    ]
    for r in trivy_results:
        p = tf.add_paragraph()
        p.text = "  🛡 " + r
        p.font.size = Pt(11)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(2)

    # ==================== SLIDE 8: ANALYSIS & REFLECTION ====================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Engineering Analysis, Reflection & Challenges", "Learnings & Interpretations")

    left = s8.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf = left.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Key Pipeline Performance Findings:"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_PRIMARY

    p_items = [
        "CI/CD Build Speed: Full pipeline completes in under 2 minutes when leveraging Docker layer caching.",
        "Shift-Left Efficiency: Security vulnerabilities caught in CI stage rather than during production penetration tests.",
        "Automated Quality Enforcement: 19 test cases act as an immutable barrier preventing regression bugs.",
        "Resource Utilization: Slim containers keep memory footprint under 90MB on EC2."
    ]
    for it in p_items:
        p = tf.add_paragraph()
        p.text = "• " + it
        p.font.size = Pt(11)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(4)

    right = s8.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.6), Inches(5.0))
    tf = right.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Challenges Overcome & Personal Reflection:"
    p.font.bold = True
    p.font.size = Pt(14)
    p.font.color.rgb = C_DARK

    c_items = [
        "Native Module Compilation in Docker: Solved `better-sqlite3` build in slim Alpine/Debian by adding multi-stage compilation dependencies.",
        "Windows Asynchronous File Locks: Debugged and resolved temp directory EPERM cleanup locks in Node.js test runners.",
        "Balancing Anonymity with Clinical Safety: Designed rule-based non-diagnostic triggers that signpost helplines without creating deanonymizing audit trails.",
        "DevSecOps Mindset: Learned that security is not a gatekeeper at the end of the waterfall, but an automated quality metric built into every commit."
    ]
    for it in c_items:
        p = tf.add_paragraph()
        p.text = "✔ " + it
        p.font.size = Pt(11)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(4)

    # ==================== SLIDE 9: CONCLUSION & FUTURE SCOPE ====================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Conclusion & Future Scope", "Looking Forward")

    box = s9.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Project Summary & Achievements:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_PRIMARY

    achievements = [
        "Successfully developed and deployed an end-to-end cloud-native wellness platform adhering to UN SDGs 3 and 9.",
        "Delivered full automation covering GitHub version control, Jenkins CI/CD, Docker containerization, and AWS Terraform IaC.",
        "Demonstrated real-world DevSecOps integration: Static Application Security Testing (SonarQube) and image vulnerability scanning (Trivy).",
        "Met 100% of the SAKEC / Mumbai University DevSecOps laboratory syllabus requirements (CO1, CO2, CO3)."
    ]
    for ac in achievements:
        p = tf.add_paragraph()
        p.text = "✔ " + ac
        p.font.size = Pt(12)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(4)

    p = tf.add_paragraph()
    p.text = "Future Scope & Enhancements:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_ACCENT
    p.space_before = Pt(12)

    future = [
        "Kubernetes Orchestration: Migrate Docker container to Amazon EKS with auto-scaling and zero-downtime rolling updates.",
        "Dynamic Application Security Testing (DAST): Embed OWASP ZAP into staging pipelines to detect runtime flaws.",
        "Cloud Secret Management: Store AES keys in AWS Secrets Manager / KMS with automated 90-day rotation.",
        "Federated OpenID Connect: Support institutional student SSO without storing campus email in application database."
    ]
    for fu in future:
        p = tf.add_paragraph()
        p.text = "🚀 " + fu
        p.font.size = Pt(12)
        p.font.color.rgb = C_MUTED
        p.space_before = Pt(4)

    # ==================== SLIDE 10: EVIDENCE OF COMPLETION ====================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Evidence of Completion & Submission Links", "Deliverables & Artifacts")

    box = s10.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0))
    tf = box.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Project Submission Deliverables:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = C_PRIMARY

    evidence = [
        ("GitHub Repository", "Contains full application source code, Dockerfile, Jenkinsfile, Terraform IaC, Ansible, and Prometheus configs."),
        ("Demonstration Video (LinkedIn)", "Full 4-minute video demonstrating working website, Docker containerization, Jenkins pipeline, Terraform, and Trivy scan."),
        ("Comprehensive Academic Report", "35+ page detailed technical documentation covering architecture, threat modeling, vulnerability tables, and reflection."),
        ("Security Audit Reports", "SonarQube SAST analysis (Quality Gate A) and Trivy HTML vulnerability scan artifacts."),
        ("Live Demonstration", "Accessible on EC2 host / local Docker environment on port 3000.")
    ]
    for label, desc in evidence:
        p = tf.add_paragraph()
        p.text = f"🔗 {label}: "
        p.font.bold = True
        p.font.size = Pt(12)
        p.font.color.rgb = C_DARK
        p.space_before = Pt(6)

        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.size = Pt(11)
        run.font.color.rgb = C_MUTED

    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    out = os.path.join(r"d:\easeway", "DEVSECOPS_PRESENTATION.pptx")
    create_presentation(out)
